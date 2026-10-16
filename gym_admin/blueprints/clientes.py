from flask import (
    Blueprint,
    request,
    render_template,
    redirect,
    url_for,
    flash,
    jsonify,
    session,
    g,
    current_app,
    Response,
)
from ..extensions import db
from sqlalchemy import text
from io import BytesIO, StringIO
import csv
from ..models import Cliente, Sucursal
from ..services.client_service import (
    client_search_filter,
    create_client,
    update_client,
    delete_client,
    compute_age,
)
from ..services.membership_service import compute_coverage, local_today
from flask_login import current_user
from ..utils import (
    admin_required,
    api_key_required,
    verify_sucursal_access,
)
from datetime import datetime, timedelta
from ..utils_timezone import utcnow

clientes_bp = Blueprint("clientes", __name__)


def _compute_client_flags_global(
    clients_list, horizon_days=7, sucursal_id=None, include_manual=False
):
    """Flags de membresía por cliente para listados y filtros.

    Usa compute_coverage: el cliente está cubierto hasta la fecha_fin más
    lejana entre sus membresías activas o en espera pagadas, así que una
    renovación ya pagada evita que figure como "por vencer".
    Returns dict cliente.id -> flags.
    """
    from ..services.attendance_service import manual_attendance_summary

    today = local_today()
    flags = {}
    for c in clients_list:
        membresias = c.membresias
        if sucursal_id:
            membresias = [m for m in membresias if m.sucursal_id == sucursal_id]
        try:
            coverage = compute_coverage(membresias, horizon_days, today=today)
        except Exception:
            current_app.logger.exception("Error calculando cobertura del cliente %s", c.id)
            coverage = {"has": False, "soon": False, "coverage_end": None}

        entry = {
            "has_active": coverage["has"],
            "soon": coverage["soon"],
            "fecha_fin": str(coverage["coverage_end"]) if coverage["soon"] else None,
        }
        if include_manual:
            entry["has_manual"] = False
            entry["manual_membresia_id"] = None
            for m in membresias:
                try:
                    if m.is_manual_asistencia and manual_attendance_summary(m).get(
                        "puede_marcar"
                    ):
                        entry["has_manual"] = True
                        entry["manual_membresia_id"] = m.id
                except Exception:
                    pass
        flags[c.id] = entry
    return flags


@clientes_bp.route("/")
def list_clientes():
    # filtros
    q = request.args.get("q", type=str)
    estado = request.args.get("estado")
    tipo = request.args.get("tipo")

    # --- LÓGICA DE FILTRADO POR SUCURSAL UNIFICADA ---
    sucursal_id = None
    if current_user and current_user.is_authenticated:
        if current_user.sucursal_id:
            sucursal_id = current_user.sucursal_id
        elif "sucursal_id" in session and session.get("sucursal_id") != 0:
            # Usar la sucursal de la sesión solo si el usuario no tiene una asignada
            sucursal_id = int(session.get("sucursal_id"))

    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 10, type=int)

    # support `expiring` filter: number of days from today to consider memberships that will expire
    expiring_days = request.args.get("expiring", type=int)
    expiring_client_ids = None
    if expiring_days:
        try:
            # Build a base client query applying the same search/estado/tipo/sucursal filters
            base_q = Cliente.query
            if q:
                base_q = base_q.filter(
                    client_search_filter(q)
                )
            if estado:
                base_q = base_q.filter_by(estado=estado)
            if tipo:
                base_q = base_q.filter_by(tipo_cliente=tipo)
            if sucursal_id and sucursal_id > 0:
                base_q = base_q.filter_by(sucursal_id=sucursal_id)

            # Fetch candidate clients (we need their memberships to compute flags)
            candidate_clients = base_q.all()
            # compute flags using the same business rules, respetando la sucursal
            flags = _compute_client_flags_global(
                candidate_clients, horizon_days=expiring_days, sucursal_id=sucursal_id
            )
            expiring_client_ids = [cid for cid, f in flags.items() if f.get("soon")]
        except Exception:
            expiring_client_ids = None

    query = Cliente.query
    if q:
        query = query.filter(
            client_search_filter(q)
        )
    if estado:
        query = query.filter_by(estado=estado)
    if tipo:
        query = query.filter_by(tipo_cliente=tipo)
    # Si sucursal_id es None -> mostrar todos (incluyendo NULL), si tiene valor >0 -> filtrar
    if sucursal_id and sucursal_id > 0:
        query = query.filter_by(sucursal_id=sucursal_id)

    # apply expiring client filter if requested
    if expiring_days is not None:
        if expiring_client_ids is None:
            # something went wrong computing ids; return no results
            query = Cliente.query.filter(Cliente.id == None)
        else:
            if len(expiring_client_ids) == 0:
                query = Cliente.query.filter(Cliente.id == None)
            else:
                query = query.filter(Cliente.id.in_(expiring_client_ids))

    pag = query.order_by(Cliente.nombre.asc(), Cliente.apellido.asc()).paginate(
        page=page, per_page=per_page, error_out=False
    )

    items = pag.items
    client_flags = _compute_client_flags_global(items, include_manual=True)

    # pass current filter params so the template can reflect selected values
    return render_template(
        "clientes/list.html",
        pagination=pag,
        clientes=items,
        today=local_today(),
        reminders=client_flags,
        q=q,
        estado=estado,
        tipo=tipo,
        per_page=per_page,
        page=page,
        expiring_days=expiring_days,
    )


@clientes_bp.route("/<int:cliente_id>")
def view_cliente(cliente_id):
    # cargar cliente con relaciones relevantes para detalle
    c = Cliente.query.get_or_404(cliente_id)
    # force load relaciones (membresias, pagos, historiales) para evitar consultas en la plantilla
    membresias = []
    try:
        for m in c.membresias:
            pagos = []
            try:
                for p in m.pagos:
                    historial = list(p.historial)
                    pagos.append({"pago": p, "historial": historial})
            except Exception:
                pagos = []
            membresias.append({"membresia": m, "pagos": pagos})
    except Exception:
        membresias = []

    return render_template(
        "clientes/detail.html", cliente=c, membresias=membresias, today=local_today()
    )


@clientes_bp.route("/search")
def search_clientes():
    # endpoint JSON para búsqueda en tiempo real
    q = request.args.get("q", type=str) or ""
    if not q:
        return jsonify([])
    query = Cliente.query.filter(
        client_search_filter(q)
    )
    # --- FILTRAR BÚSQUEDA POR SUCURSAL ---
    sucursal_id = None
    if current_user and current_user.is_authenticated:
        if current_user.sucursal_id:
            sucursal_id = current_user.sucursal_id
        elif "sucursal_id" in session and session.get("sucursal_id") != 0:
            sucursal_id = session.get("sucursal_id")

    if sucursal_id:
        query = query.filter(Cliente.sucursal_id == sucursal_id)

    rows = query.limit(50).all()
    result = [
        {"id": r.id, "text": f"{r.nombre} {r.apellido} - {r.correo or r.telefono}"}
        for r in rows
    ]
    return jsonify(result)


@clientes_bp.route("/table")
def clientes_table():
    # devuelve la tabla renderizada (partial) para recarga por AJAX
    # Reusar la misma lógica de filtros que list_clientes
    q = request.args.get("q", type=str)
    estado = request.args.get("estado")
    tipo = request.args.get("tipo")
    # page y per_page son incluidos para la renderización
    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 10, type=int)

    # --- LÓGICA DE FILTRADO POR SUCURSAL UNIFICADA (igual que list_clientes) ---
    sucursal_id = None
    if current_user and current_user.is_authenticated:
        if current_user.sucursal_id:
            sucursal_id = current_user.sucursal_id
        elif "sucursal_id" in session and session.get("sucursal_id") != 0:
            # Usar la sucursal de la sesión solo si el usuario no tiene una asignada
            sucursal_id = int(session.get("sucursal_id"))

    query = Cliente.query
    if q:
        query = query.filter(
            client_search_filter(q)
        )
    if estado:
        query = query.filter_by(estado=estado)
    if tipo:
        query = query.filter_by(tipo_cliente=tipo)
    if sucursal_id and sucursal_id > 0:
        query = query.filter_by(sucursal_id=sucursal_id)

    # support expiring filter for AJAX partials
    expiring_days = request.args.get("expiring", type=int)
    if expiring_days:
        try:
            # Build base query applying same filters
            base_q = Cliente.query
            if q:
                base_q = base_q.filter(
                    client_search_filter(q)
                )
            if estado:
                base_q = base_q.filter_by(estado=estado)
            if tipo:
                base_q = base_q.filter_by(tipo_cliente=tipo)
            if sucursal_id and sucursal_id > 0:
                base_q = base_q.filter_by(sucursal_id=sucursal_id)
            candidate_clients = base_q.all()
            flags = _compute_client_flags_global(
                candidate_clients, horizon_days=expiring_days, sucursal_id=sucursal_id
            )
            expiring_client_ids = [cid for cid, f in flags.items() if f.get("soon")]
            if expiring_client_ids:
                query = query.filter(Cliente.id.in_(expiring_client_ids))
            else:
                query = Cliente.query.filter(Cliente.id == None)
        except Exception:
            query = Cliente.query.filter(Cliente.id == None)

    pag = query.order_by(Cliente.nombre.asc(), Cliente.apellido.asc()).paginate(
        page=page, per_page=per_page, error_out=False
    )

    mode = request.args.get("mode")


    client_flags = _compute_client_flags_global(pag.items, include_manual=True)

    # include total count for UI (pagination.total)
    ctx = dict(
        pagination=pag,
        clientes=pag.items,
        today=local_today(),
        reminders=client_flags,
        total=pag.total,
        q=q,
        estado=estado,
        tipo=tipo,
        per_page=per_page,
        page=page,
        expiring_days=expiring_days,
    )
    if mode == "accordion":
        return render_template("clientes/_accordion.html", **ctx)
    return render_template("clientes/_table.html", **ctx)


@clientes_bp.route("/export.csv")
def export_clientes_csv():
    # Exportar clientes actuales según filtros a CSV
    q = request.args.get("q", type=str)
    estado = request.args.get("estado")
    tipo = request.args.get("tipo")
    sucursal_id = request.args.get("sucursal_id", type=int)

    query = Cliente.query
    if q:
        query = query.filter(
            client_search_filter(q)
        )
    if estado:
        query = query.filter_by(estado=estado)
    if tipo:
        query = query.filter_by(tipo_cliente=tipo)
    if sucursal_id and sucursal_id > 0:
        query = query.filter_by(sucursal_id=sucursal_id)

    rows = query.order_by(Cliente.nombre.asc(), Cliente.apellido.asc()).all()

    import csv
    from io import StringIO

    si = StringIO()
    cw = csv.writer(si)
    cw.writerow(
        [
            "id",
            "nombre",
            "apellido",
            "correo",
            "telefono",
            "fecha_nacimiento",
            "tipo_cliente",
            "estado",
            "sucursal_id",
        ]
    )
    for r in rows:
        cw.writerow(
            [
                r.id,
                r.nombre,
                r.apellido,
                r.correo or "",
                r.telefono or "",
                r.fecha_nacimiento or "",
                r.tipo_cliente,
                r.estado,
                r.sucursal_id or "",
            ]
        )
    output = si.getvalue()
    return (
        output,
        200,
        {
            "Content-Type": "text/csv",
            "Content-Disposition": 'attachment; filename="clientes_export.csv"',
        },
    )


@clientes_bp.route("/print")
def print_clientes():
    # vista optimizada para imprimir según filtros
    q = request.args.get("q", type=str)
    estado = request.args.get("estado")
    tipo = request.args.get("tipo")
    sucursal_id = request.args.get("sucursal_id", type=int)

    query = Cliente.query
    if q:
        query = query.filter(
            client_search_filter(q)
        )
    if estado:
        query = query.filter_by(estado=estado)
    if tipo:
        query = query.filter_by(tipo_cliente=tipo)
    if sucursal_id and sucursal_id > 0:
        query = query.filter_by(sucursal_id=sucursal_id)

    rows = query.order_by(Cliente.nombre.asc(), Cliente.apellido.asc()).all()

    return render_template("clientes/print.html", clientes=rows, now=utcnow())


@clientes_bp.route("/create", methods=["GET", "POST"])
def create_cliente():
    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"

    if request.method == "POST":
        nombre = (request.form.get("nombre") or "").strip()
        apellido = (request.form.get("apellido") or "").strip()
        correo = (request.form.get("correo") or "").strip().lower()
        telefono = (request.form.get("telefono") or "").strip()
        fecha_nacimiento_raw = request.form.get("fecha_nacimiento")
        tipo_cliente = request.form.get("tipo_cliente") or "Individual"

        # Validar fecha y edad antes de llamar al servicio
        errors = {}
        if not fecha_nacimiento_raw:
            errors["fecha_nacimiento"] = "Fecha de nacimiento requerida"
        else:
            try:
                fecha_nacimiento = datetime.strptime(
                    fecha_nacimiento_raw, "%Y-%m-%d"
                ).date()
            except Exception:
                errors["fecha_nacimiento"] = "Fecha de nacimiento inválida"
                fecha_nacimiento = None

            if fecha_nacimiento:
                edad = compute_age(fecha_nacimiento)
                if edad is None or edad < 12:
                    errors["fecha_nacimiento"] = "La edad del cliente debe ser al menos 12 años"

        if errors:
            if is_ajax:
                return jsonify({"success": False, "message": "Error de validación", "errors": errors}), 400
            else:
                for field, msg in errors.items():
                    flash(msg)
                return redirect(url_for("clientes.create_cliente"))

        # Determinar sucursal
        sucursal_id = None
        if current_user and current_user.is_authenticated and current_user.sucursal_id:
            sucursal_id = current_user.sucursal_id
        elif request.form.get("sucursal_id"):
            try:
                sucursal_id = int(request.form.get("sucursal_id"))
            except (ValueError, TypeError):
                pass
        elif session.get("sucursal_id") and str(session.get("sucursal_id")) != "0":
            try:
                sucursal_id = int(session.get("sucursal_id"))
            except (ValueError, TypeError):
                pass

        # Usar servicio para crear cliente
        success, cliente, error_msg = create_client(
            nombre=nombre,
            apellido=apellido,
            correo=correo or None,
            telefono=telefono or None,
            fecha_nacimiento=fecha_nacimiento,
            tipo_cliente=tipo_cliente,
            sucursal_id=sucursal_id,
        )

        if success:
            if is_ajax:
                return jsonify({
                    "success": True,
                    "message": "Cliente creado exitosamente",
                    "cliente_id": cliente.id
                })
            else:
                flash("Cliente creado")
                return redirect(url_for("clientes.list_clientes"))
        else:
            if is_ajax:
                return jsonify({"success": False, "message": f"Error: {error_msg}"}), 400
            else:
                flash(f"Error: {error_msg}")
                return redirect(url_for("clientes.create_cliente"))

    # GET - Cargar formulario vacío
    sucursales = Sucursal.query.order_by(Sucursal.nombre).all()
    sucursal_actual_id = None
    if current_user and current_user.is_authenticated:
        if current_user.sucursal_id:
            sucursal_actual_id = current_user.sucursal_id
        elif session.get("sucursal_id") and str(session.get("sucursal_id")) != "0":
            sucursal_actual_id = int(session.get("sucursal_id"))

    # Si es AJAX, devolver solo el formulario sin estructura de página
    if is_ajax:
        return render_template(
            "clientes/form_modal.html",
            cliente=None,
            sucursales=sucursales,
            sucursal_actual_id=sucursal_actual_id,
        )

    # Si es página regular, devolver con estructura completa
    return render_template(
        "clientes/form.html",
        cliente=None,
        sucursales=sucursales,
        sucursal_actual_id=sucursal_actual_id,
    )


@clientes_bp.route("/<int:cliente_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_cliente(cliente_id):
    c = Cliente.query.get_or_404(cliente_id)
    verify_sucursal_access(c)
    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"

    if request.method == "POST":
        nombre = (request.form.get("nombre") or "").strip()
        apellido = (request.form.get("apellido") or "").strip()
        correo = (request.form.get("correo") or "").strip().lower()
        telefono = (request.form.get("telefono") or "").strip()
        fecha_nacimiento_raw = request.form.get("fecha_nacimiento")
        tipo_cliente = request.form.get("tipo_cliente") or c.tipo_cliente

        # Validar edad
        errors = {}
        if fecha_nacimiento_raw:
            try:
                fecha_nacimiento = datetime.strptime(
                    fecha_nacimiento_raw, "%Y-%m-%d"
                ).date()
            except Exception:
                errors["fecha_nacimiento"] = "Fecha de nacimiento inválida"
                fecha_nacimiento = c.fecha_nacimiento

            if not errors and fecha_nacimiento:
                edad = compute_age(fecha_nacimiento)
                if edad is None or edad < 12:
                    errors["fecha_nacimiento"] = "La edad del cliente debe ser al menos 12 años"
        else:
            fecha_nacimiento = c.fecha_nacimiento

        if errors:
            if is_ajax:
                return jsonify({"success": False, "message": "Error de validación", "errors": errors}), 400
            else:
                for field, msg in errors.items():
                    flash(msg)
                return redirect(url_for("clientes.edit_cliente", cliente_id=cliente_id))

        # Determinar sucursal
        sucursal_id = c.sucursal_id
        if request.form.get("sucursal_id"):
            try:
                sucursal_id = int(request.form.get("sucursal_id"))
            except (ValueError, TypeError):
                pass

        # Usar servicio para actualizar
        success, error_msg = update_client(
            cliente_id=cliente_id,
            nombre=nombre,
            apellido=apellido,
            correo=correo or None,
            telefono=telefono or None,
            fecha_nacimiento=fecha_nacimiento,
            tipo_cliente=tipo_cliente,
            sucursal_id=sucursal_id,
        )

        if success:
            if is_ajax:
                return jsonify({
                    "success": True,
                    "message": "Cliente actualizado exitosamente",
                    "cliente_id": cliente_id
                })
            else:
                flash("Cliente actualizado")
                return redirect(url_for("clientes.list_clientes"))
        else:
            if is_ajax:
                return jsonify({"success": False, "message": f"Error: {error_msg}"}), 400
            else:
                flash(f"Error: {error_msg}")
                return redirect(url_for("clientes.edit_cliente", cliente_id=cliente_id))

    # GET - Cargar formulario con datos existentes
    if is_ajax:
        sucursales = Sucursal.query.order_by(Sucursal.nombre).all()
        return render_template(
            "clientes/form_modal.html",
            cliente=c,
            sucursales=sucursales,
            sucursal_actual_id=c.sucursal_id
        )

    return render_template("clientes/form.html", cliente=c)


@clientes_bp.route("/<int:cliente_id>/delete", methods=["POST"])
@admin_required
def delete_cliente(cliente_id):
    c = Cliente.query.get_or_404(cliente_id)
    verify_sucursal_access(c)
    try:
        pass
    except Exception:
        pass

    if not (
        current_user
        and getattr(current_user, "is_authenticated", False)
        and getattr(current_user, "rol_id", None) == 1
    ):
        flash("No autorizado")
        return redirect(url_for("index"))

    success, error_msg = delete_client(cliente_id)
    if success:
        flash("Cliente eliminado")
    else:
        flash(f"Error: {error_msg}")

    return redirect(url_for("clientes.list_clientes"))


@clientes_bp.route("/api/list", methods=["GET"])
@api_key_required  # Este decorador debe setear g.api_key_user y g.api_key_sucursal_id
def api_list_clientes():
    try:
        # 1. Obtener contexto de seguridad desde el decorador
        getattr(g, "api_key_user", None)
        api_sucursal_id = getattr(g, "api_key_sucursal_id", None)

        # 2. Obtener y sanear parámetros de búsqueda
        estado = request.args.get("estado", type=str)
        page = max(1, request.args.get("page", 1, type=int))
        per_page = min(100, max(1, request.args.get("per_page", 50, type=int)))
        vencer_en_dias = request.args.get("vencer_en_dias", type=int)

        # 3. Lógica de Sucursal (Prioridad: API Key > Query Param)
        sucursal_id = api_sucursal_id or request.args.get("sucursal_id", type=int)

        # 4. Construcción de Query Base
        query = Cliente.query

        if estado:
            if estado not in ["Activo", "Inactivo"]:
                return (
                    jsonify(
                        {
                            "success": False,
                            "error": 'Estado inválido. Debe ser "Activo" o "Inactivo"',
                            "data": None,
                        }
                    ),
                    400,
                )
            query = query.filter_by(estado=estado)

        if sucursal_id:
            query = query.filter_by(sucursal_id=sucursal_id)

        # 5. Filtro de Vencimiento (Optimizado)
        if vencer_en_dias is not None:
            vencer_en_dias = max(1, vencer_en_dias)
            hoy = local_today()
            horizonte = hoy + timedelta(days=vencer_en_dias)

            # Filtramos clientes que tengan membresías activas o en espera
            # cuya fecha_fin esté entre hoy y el horizonte.
            from ..models import EstadoMembresia
            activa_estado = EstadoMembresia.query.filter(EstadoMembresia.nombre.ilike("activ%")).first()
            espera_estado = EstadoMembresia.query.filter(EstadoMembresia.nombre.ilike("%espera%")).first()
            activa_id = activa_estado.id if activa_estado else 2
            espera_id = espera_estado.id if espera_estado else 5

            query = (
                query.join(Cliente.membresias)
                .filter(
                    Cliente.membresias.property.mapper.class_.estado_id.in_([activa_id, espera_id]),
                    (Cliente.membresias.property.mapper.class_.fecha_fin >= hoy),
                    (Cliente.membresias.property.mapper.class_.fecha_fin <= horizonte),
                )
                .distinct()
            )

        # 6. Ejecutar Paginación
        pagination = query.order_by(
            Cliente.nombre.asc(), Cliente.apellido.asc()
        ).paginate(page=page, per_page=per_page, error_out=False)

        # 7. Serialización de Datos
        clientes_data = []
        for cliente in pagination.items:
            # Lógica de membresía para el JSON de respuesta
            prox_venc = None
            tiene_activa = False

            # Buscamos la membresía más cercana a vencer
            membresias_validas = [
                m
                for m in cliente.membresias
                if getattr(m, "estado_id", None) == 2 and m.fecha_fin
            ]
            if membresias_validas:
                # Ordenamos por fecha_fin para obtener la más próxima
                m_proxima = min(membresias_validas, key=lambda x: x.fecha_fin)
                if m_proxima.fecha_fin >= local_today():
                    prox_venc = m_proxima.fecha_fin
                    tiene_activa = True

            clientes_data.append(
                {
                    "id": cliente.id,
                    "nombre": cliente.nombre or "",
                    "apellido": cliente.apellido or "",
                    "correo": cliente.correo or "",
                    "telefono": cliente.telefono or "",
                    "fecha_nacimiento": (
                        str(cliente.fecha_nacimiento)
                        if cliente.fecha_nacimiento
                        else None
                    ),
                    "estado": cliente.estado,
                    "sucursal_id": cliente.sucursal_id,
                    "tiene_membresia_activa": tiene_activa,
                    "fecha_vencimiento_proxima": str(prox_venc) if prox_venc else None,
                    "dias_para_vencer": (
                        (prox_venc - local_today()).days if prox_venc else None
                    ),
                }
            )

        return (
            jsonify(
                {
                    "success": True,
                    "data": {
                        "clientes": clientes_data,
                        "pagination": {
                            "total": pagination.total,
                            "total_pages": pagination.pages,
                            "current_page": page,
                            "per_page": per_page,
                        },
                        "filters": {
                            "estado": estado,
                            "sucursal_id": sucursal_id,
                            "vencer_en_dias": vencer_en_dias,
                        },
                    },
                }
            ),
            200,
        )

    except Exception as e:
        # Registro del error real para el desarrollador
        print(f"DEBUG ERROR: {str(e)}")
        return jsonify({"success": False, "error": "Error interno del servidor"}), 500


@clientes_bp.route("/api/reporte-consolidado", methods=["GET"])
@clientes_bp.route("/exportar-consolidado", methods=["GET"])
@api_key_required
def api_reporte_consolidado_bi():
    """
    Super-endpoint de Inteligencia de Negocios (Power BI / Excel / Google Sheets).
    Parámetros opcionales:
      - formato: 'xlsx' / 'excel', 'csv', 'json' (default 'json')
      - sucursal_id: int (filtra por sucursal si la API Key o sesión es global)
      - estado: 'Activo' | 'Inactivo'
      - q: str (búsqueda opcional por nombre, apellido, correo o teléfono)
    """
    try:
        # 1. Contexto de seguridad y filtros opcionales
        api_sucursal_id = getattr(g, "api_key_sucursal_id", None)
        req_sucursal_id = request.args.get("sucursal_id", type=int)

        # Si el usuario está autenticado por sesión web y no especificó sucursal por query
        session_suc = session.get("sucursal_id") if hasattr(session, "get") else None
        if session_suc in (0, "0", ""):
            session_suc = None
        session_sucursal_id = int(session_suc) if session_suc else None

        sucursal_id = api_sucursal_id or req_sucursal_id or session_sucursal_id
        estado = request.args.get("estado", type=str)
        if estado:
            estado = estado.strip()
        q = request.args.get("q", type=str)
        formato = (request.args.get("formato", "json") or "json").lower().strip()

        # 2. Aumentar group_concat_max_len en la sesión actual para evitar truncamientos
        # en clientes con historiales muy largos de membresías (solo MySQL)
        try:
            if db.engine.dialect.name == "mysql":
                db.session.execute(text("SET SESSION group_concat_max_len = 100000;"))
        except Exception as dialect_err:
            current_app.logger.warning("No se pudo configurar group_concat_max_len: %s", str(dialect_err))

        # 3. Construir filtros dinámicos seguros (bindparams)
        filtros_sql = ["1 = 1"]
        params = {}

        if sucursal_id and sucursal_id > 0:
            filtros_sql.append("c.sucursal_id = :sucursal_id")
            params["sucursal_id"] = sucursal_id

        if estado in ("Activo", "Inactivo"):
            filtros_sql.append("c.estado = :estado")
            params["estado"] = estado

        if q and q.strip():
            filtros_sql.append(
                "(c.nombre LIKE :q OR c.apellido LIKE :q OR c.correo LIKE :q OR c.telefono LIKE :q)"
            )
            params["q"] = f"%{q.strip()}%"

        where_clause = " AND ".join(filtros_sql)

        # 4. Superconsulta SQL Consolidada
        sql_query = text(f"""
            SELECT 
                c.id AS cliente_id,
                c.nombre,
                c.apellido,
                TRIM(CONCAT(COALESCE(c.nombre, ''), ' ', COALESCE(c.apellido, ''))) AS nombre_completo,
                COALESCE(c.telefono, '') AS telefono,
                COALESCE(c.correo, '') AS correo,
                c.tipo_cliente,
                c.estado AS estado_cliente,
                c.sucursal_id,
                COALESCE(s.nombre, 'Sin Sucursal') AS sucursal_nombre,
                DATE_FORMAT(c.fecha_nacimiento, '%Y-%m-%d') AS fecha_nacimiento,
                TIMESTAMPDIFF(YEAR, c.fecha_nacimiento, CURDATE()) AS edad_actual,
                CASE 
                    WHEN c.fecha_nacimiento IS NULL THEN 'No registrado'
                    WHEN TIMESTAMPDIFF(YEAR, c.fecha_nacimiento, CURDATE()) < 18 THEN 'Menor de 18'
                    WHEN TIMESTAMPDIFF(YEAR, c.fecha_nacimiento, CURDATE()) BETWEEN 18 AND 24 THEN '18-24 años'
                    WHEN TIMESTAMPDIFF(YEAR, c.fecha_nacimiento, CURDATE()) BETWEEN 25 AND 34 THEN '25-34 años'
                    WHEN TIMESTAMPDIFF(YEAR, c.fecha_nacimiento, CURDATE()) BETWEEN 35 AND 44 THEN '35-44 años'
                    WHEN TIMESTAMPDIFF(YEAR, c.fecha_nacimiento, CURDATE()) BETWEEN 45 AND 54 THEN '45-54 años'
                    ELSE '55+ años'
                END AS rango_etario,

                DATE_FORMAT(c.fecha_registro, '%Y-%m-%d') AS fecha_registro,
                DATEDIFF(CURDATE(), c.fecha_registro) AS dias_desde_registro,
                CASE 
                    WHEN rm.primera_fecha_inicio IS NOT NULL AND c.fecha_registro IS NOT NULL 
                    THEN GREATEST(0, DATEDIFF(rm.primera_fecha_inicio, c.fecha_registro))
                    ELSE NULL 
                END AS dias_en_activar_primera_membresia,

                COALESCE(rm.total_membresias_historial, 0) AS total_membresias_solicitadas,
                COALESCE(rm.membresias_activadas_pagadas, 0) AS membresias_activadas_pagadas,
                GREATEST(0, COALESCE(rm.membresias_activadas_pagadas, 0) - 1) AS cantidad_renovaciones,
                COALESCE(rm.renovaciones_tipo_extend, 0) AS renovaciones_marcadas_extend,
                CASE 
                    WHEN COALESCE(rm.membresias_activadas_pagadas, 0) > 1 THEN 'Recurrente'
                    WHEN COALESCE(rm.membresias_activadas_pagadas, 0) = 1 THEN 'Única Membresía'
                    ELSE 'Sin Membresías Pagadas'
                END AS tipo_fidelidad_cliente,

                COALESCE(rm.planes_cortos_dias_semanas, 0) AS membresias_cortas_diarias_semanales,
                COALESCE(rm.planes_1_mes, 0) AS membresias_1_mes,
                COALESCE(rm.planes_2_meses, 0) AS membresias_2_meses,
                COALESCE(rm.planes_3_meses, 0) AS membresias_3_meses,
                COALESCE(rm.planes_6_meses, 0) AS membresias_6_meses,
                COALESCE(rm.planes_12_meses_anual, 0) AS membresias_anuales,
                COALESCE(rm.planes_otros_meses, 0) AS membresias_otros_plazos,
                COALESCE(rm.planes_por_cupo_sesiones, 0) AS membresias_por_cupo_asistencia,

                COALESCE(rm.periodo_primera_membresia, 'Sin membresías') AS periodo_primera_membresia,
                DATE_FORMAT(rm.primera_fecha_inicio, '%Y-%m-%d') AS primera_membresia_inicio,
                rm.primera_fecha_fin AS primera_membresia_fin,
                COALESCE(rm.primer_plan_nombre, 'N/A') AS primer_plan_contratado,

                COALESCE(rm.periodo_ultima_membresia, 'Sin membresías') AS periodo_ultima_membresia,
                rm.ultima_fecha_inicio AS ultima_membresia_inicio,
                DATE_FORMAT(rm.ultima_fecha_fin, '%Y-%m-%d') AS ultima_membresia_fin,
                COALESCE(rm.ultimo_plan_nombre, 'N/A') AS ultimo_plan_contratado,
                COALESCE(rm.ultimo_estado_membresia, 'Sin Membresía') AS estado_ultima_membresia,

                CASE 
                    WHEN rm.max_fecha_fin_pagada >= CURDATE() THEN 1 
                    ELSE 0 
                END AS tiene_membresia_vigente_hoy,
                CASE 
                    WHEN rm.max_fecha_fin_pagada IS NOT NULL 
                    THEN DATEDIFF(rm.max_fecha_fin_pagada, CURDATE()) 
                    ELSE NULL 
                END AS dias_para_vencer_o_vencido,
                CASE 
                    WHEN rm.max_fecha_fin_pagada IS NOT NULL AND rm.max_fecha_fin_pagada < CURDATE() 
                    THEN DATEDIFF(CURDATE(), rm.max_fecha_fin_pagada)
                    ELSE 0 
                END AS dias_abandono_desde_ultimo_vencimiento,
                COALESCE(rm.total_dias_contratados, 0) AS total_dias_contratados_historico,
                CASE
                    WHEN COALESCE(rm.membresias_activadas_pagadas, 0) = 0 THEN '1. Prospecto (Sin Membresía Pagada)'
                    WHEN rm.max_fecha_fin_pagada >= CURDATE() AND DATEDIFF(rm.max_fecha_fin_pagada, CURDATE()) <= 7 THEN '2. Alerta: Por Vencer (<= 7 días)'
                    WHEN rm.max_fecha_fin_pagada >= CURDATE() AND COALESCE(rm.membresias_activadas_pagadas, 0) = 1 THEN '3. Activo Nuevo (1ra Membresía)'
                    WHEN rm.max_fecha_fin_pagada >= CURDATE() AND COALESCE(rm.membresias_activadas_pagadas, 0) > 1 THEN '4. Activo Fidelizado (Renovador)'
                    WHEN DATEDIFF(CURDATE(), rm.max_fecha_fin_pagada) BETWEEN 1 AND 30 THEN '5. En Riesgo de Abandono (Vencido 1-30 días)'
                    WHEN DATEDIFF(CURDATE(), rm.max_fecha_fin_pagada) BETWEEN 31 AND 90 THEN '6. Abandono Reciente (Vencido 31-90 días)'
                    ELSE '7. Cliente Perdido / Churn (> 90 días)'
                END AS segmento_fidelidad_churn,

                ROUND(COALESCE(rm.total_pagado_membresias, 0), 2) AS total_pagado_membresias,
                ROUND(
                    CASE 
                        WHEN COALESCE(rm.membresias_activadas_pagadas, 0) > 0 
                        THEN COALESCE(rm.total_pagado_membresias, 0) / rm.membresias_activadas_pagadas 
                        ELSE 0 
                    END, 2
                ) AS ticket_promedio_membresia,
                ROUND(COALESCE(rm.saldo_pendiente_total, 0), 2) AS saldo_pendiente_membresias,
                COALESCE(rm.metodo_pago_frecuente, 'N/A') AS metodo_pago_preferido,
                COALESCE(rm.membresias_con_promocion, 0) AS membresias_con_promocion,
                COALESCE(rv.total_compras_pos, 0) AS cantidad_compras_productos_pos,
                ROUND(COALESCE(rv.monto_total_ventas_pos, 0), 2) AS total_gastado_productos_pos,
                ROUND(COALESCE(rm.total_pagado_membresias, 0) + COALESCE(rv.monto_total_ventas_pos, 0), 2) AS ltv_ingresos_totales_cliente,

                GREATEST(COALESCE(ra.total_marcajes_tabla, 0), COALESCE(rm.total_marcajes_usados_col, 0)) AS total_asistencias_registradas,
                DATE_FORMAT(ra.ultima_fecha_asistencia, '%Y-%m-%d %H:%i') AS ultima_fecha_asistencia,
                CASE 
                    WHEN ra.ultima_fecha_asistencia IS NOT NULL 
                    THEN DATEDIFF(CURDATE(), DATE(ra.ultima_fecha_asistencia)) 
                    ELSE NULL 
                END AS dias_sin_asistir

            FROM clientes c
            LEFT JOIN sucursales s ON s.id = c.sucursal_id
            LEFT JOIN (
                SELECT 
                    m_det.cliente_id,
                    SUM(CASE WHEN m_det.es_anulada = 0 THEN 1 ELSE 0 END) AS total_membresias_historial,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 THEN 1 ELSE 0 END) AS membresias_activadas_pagadas,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND m_det.tipo_creacion = 'extend' THEN 1 ELSE 0 END) AS renovaciones_tipo_extend,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND m_det.promocion_id IS NOT NULL THEN 1 ELSE 0 END) AS membresias_con_promocion,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND m_det.categoria_duracion = 'CORTA' THEN 1 ELSE 0 END) AS planes_cortos_dias_semanas,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND m_det.categoria_duracion = '1_MES' THEN 1 ELSE 0 END) AS planes_1_mes,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND m_det.categoria_duracion = '2_MESES' THEN 1 ELSE 0 END) AS planes_2_meses,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND m_det.categoria_duracion = '3_MESES' THEN 1 ELSE 0 END) AS planes_3_meses,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND m_det.categoria_duracion = '6_MESES' THEN 1 ELSE 0 END) AS planes_6_meses,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND m_det.categoria_duracion = '12_MESES' THEN 1 ELSE 0 END) AS planes_12_meses_anual,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND m_det.categoria_duracion = 'OTROS' THEN 1 ELSE 0 END) AS planes_otros_meses,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 AND (m_det.tipo_uso = 'asistencia_manual' OR m_det.cupo_marcajes IS NOT NULL) THEN 1 ELSE 0 END) AS planes_por_cupo_sesiones,
                    MIN(CASE WHEN m_det.es_anulada = 0 THEN m_det.fecha_inicio END) AS primera_fecha_inicio,
                    SUBSTRING_INDEX(GROUP_CONCAT(CASE WHEN m_det.es_anulada = 0 THEN DATE_FORMAT(m_det.fecha_fin, '%Y-%m-%d') END ORDER BY m_det.fecha_inicio ASC, m_det.membresia_id ASC SEPARATOR '||'), '||', 1) AS primera_fecha_fin,
                    SUBSTRING_INDEX(GROUP_CONCAT(CASE WHEN m_det.es_anulada = 0 THEN m_det.plan_nombre END ORDER BY m_det.fecha_inicio ASC, m_det.membresia_id ASC SEPARATOR '||'), '||', 1) AS primer_plan_nombre,
                    SUBSTRING_INDEX(GROUP_CONCAT(CASE WHEN m_det.es_anulada = 0 THEN CONCAT(DATE_FORMAT(m_det.fecha_inicio, '%Y-%m-%d'), ' al ', DATE_FORMAT(m_det.fecha_fin, '%Y-%m-%d'), ' (', m_det.plan_nombre, ')') END ORDER BY m_det.fecha_inicio ASC, m_det.membresia_id ASC SEPARATOR '||'), '||', 1) AS periodo_primera_membresia,
                    SUBSTRING_INDEX(GROUP_CONCAT(CASE WHEN m_det.es_anulada = 0 THEN DATE_FORMAT(m_det.fecha_inicio, '%Y-%m-%d') END ORDER BY m_det.fecha_inicio DESC, m_det.membresia_id DESC SEPARATOR '||'), '||', 1) AS ultima_fecha_inicio,
                    MAX(CASE WHEN m_det.es_anulada = 0 THEN m_det.fecha_fin END) AS ultima_fecha_fin,
                    SUBSTRING_INDEX(GROUP_CONCAT(CASE WHEN m_det.es_anulada = 0 THEN m_det.plan_nombre END ORDER BY m_det.fecha_inicio DESC, m_det.membresia_id DESC SEPARATOR '||'), '||', 1) AS ultimo_plan_nombre,
                    SUBSTRING_INDEX(GROUP_CONCAT(CASE WHEN m_det.es_anulada = 0 THEN m_det.estado_nombre END ORDER BY m_det.fecha_inicio DESC, m_det.membresia_id DESC SEPARATOR '||'), '||', 1) AS ultimo_estado_membresia,
                    SUBSTRING_INDEX(GROUP_CONCAT(CASE WHEN m_det.es_anulada = 0 THEN CONCAT(DATE_FORMAT(m_det.fecha_inicio, '%Y-%m-%d'), ' al ', DATE_FORMAT(m_det.fecha_fin, '%Y-%m-%d'), ' (', m_det.plan_nombre, ')') END ORDER BY m_det.fecha_inicio DESC, m_det.membresia_id DESC SEPARATOR '||'), '||', 1) AS periodo_ultima_membresia,
                    MAX(CASE WHEN m_det.es_pagada_y_activada = 1 THEN m_det.fecha_fin END) AS max_fecha_fin_pagada,
                    SUM(CASE WHEN m_det.es_pagada_y_activada = 1 THEN GREATEST(0, DATEDIFF(m_det.fecha_fin, m_det.fecha_inicio)) ELSE 0 END) AS total_dias_contratados,
                    SUM(CASE WHEN m_det.es_anulada = 0 THEN m_det.monto_pagado ELSE 0 END) AS total_pagado_membresias,
                    SUM(CASE WHEN m_det.es_anulada = 0 THEN m_det.saldo_pendiente ELSE 0 END) AS saldo_pendiente_total,
                    SUBSTRING_INDEX(GROUP_CONCAT(m_det.metodo_pago ORDER BY m_det.fecha_inicio DESC, m_det.membresia_id DESC SEPARATOR '||'), '||', 1) AS metodo_pago_frecuente,
                    SUM(COALESCE(m_det.marcajes_usados, 0)) AS total_marcajes_usados_col
                FROM (
                    SELECT 
                        m.id AS membresia_id,
                        m.cliente_id,
                        m.fecha_inicio,
                        m.fecha_fin,
                        m.tipo_creacion,
                        m.tipo_uso,
                        m.cupo_marcajes,
                        m.marcajes_usados,
                        m.promocion_id,
                        p.nombre AS plan_nombre,
                        COALESCE(em.nombre, 'Desconocido') AS estado_nombre,
                        CASE 
                            WHEN m.estado_id = 3 OR LOWER(COALESCE(em.nombre, '')) LIKE 'anulad%' THEN 1 
                            ELSE 0 
                        END AS es_anulada,
                        CASE 
                            WHEN (m.estado_id != 3 AND LOWER(COALESCE(em.nombre, '')) NOT LIKE 'anulad%')
                                 AND (m.estado_id != 1 AND LOWER(COALESCE(em.nombre, '')) NOT LIKE 'pendient%')
                                 AND COALESCE(pg_agg.monto_pagado_real, 0) > 0 
                            THEN 1 
                            ELSE 0 
                        END AS es_pagada_y_activada,
                        CASE 
                            WHEN LOWER(p.duracion_tipo) LIKE 'dia%' OR LOWER(p.duracion_tipo) LIKE 'd%' OR LOWER(p.duracion_tipo) LIKE 'seman%' 
                                 OR (DATEDIFF(m.fecha_fin, m.fecha_inicio) < 25) THEN 'CORTA'
                            WHEN (LOWER(p.duracion_tipo) LIKE 'mes%' AND p.duracion_valor = 1) 
                                 OR LOWER(p.duracion) LIKE '%1 mes%' OR LOWER(p.nombre) LIKE '%mensual%'
                                 OR (DATEDIFF(m.fecha_fin, m.fecha_inicio) BETWEEN 25 AND 35) THEN '1_MES'
                            WHEN (LOWER(p.duracion_tipo) LIKE 'mes%' AND p.duracion_valor = 2) 
                                 OR LOWER(p.duracion) LIKE '%2 mes%' OR LOWER(p.nombre) LIKE '%bimestral%'
                                 OR (DATEDIFF(m.fecha_fin, m.fecha_inicio) BETWEEN 55 AND 65) THEN '2_MESES'
                            WHEN (LOWER(p.duracion_tipo) LIKE 'mes%' AND p.duracion_valor = 3) 
                                 OR LOWER(p.duracion) LIKE '%3 mes%' OR LOWER(p.nombre) LIKE '%trimestral%'
                                 OR (DATEDIFF(m.fecha_fin, m.fecha_inicio) BETWEEN 85 AND 95) THEN '3_MESES'
                            WHEN (LOWER(p.duracion_tipo) LIKE 'mes%' AND p.duracion_valor = 6) 
                                 OR LOWER(p.duracion) LIKE '%6 mes%' OR LOWER(p.nombre) LIKE '%semestral%'
                                 OR (DATEDIFF(m.fecha_fin, m.fecha_inicio) BETWEEN 175 AND 190) THEN '6_MESES'
                            WHEN (LOWER(p.duracion_tipo) LIKE 'mes%' AND p.duracion_valor = 12) 
                                 OR LOWER(p.duracion_tipo) LIKE 'a%' OR LOWER(p.duracion) LIKE '%anual%' 
                                 OR LOWER(p.duracion) LIKE '%12 mes%' OR LOWER(p.nombre) LIKE '%anual%'
                                 OR (DATEDIFF(m.fecha_fin, m.fecha_inicio) >= 350) THEN '12_MESES'
                            ELSE 'OTROS'
                        END AS categoria_duracion,
                        COALESCE(pg_agg.monto_pagado_real, 0) AS monto_pagado,
                        GREATEST(0, COALESCE(pg_agg.monto_comprometido, p.precio, 0) - COALESCE(pg_agg.monto_pagado_real, 0)) AS saldo_pendiente,
                        pg_agg.ultimo_metodo_pago AS metodo_pago
                    FROM membresias m
                    INNER JOIN planes p ON p.id = m.plan_id
                    LEFT JOIN estados_membresias em ON em.id = m.estado_id
                    LEFT JOIN (
                        SELECT 
                            pg.membresia_id,
                            SUM(COALESCE(pg.monto_total, 0)) AS monto_comprometido,
                            SUM(
                                CASE 
                                    WHEN hp_sum.total_abonos IS NOT NULL AND hp_sum.total_abonos > 0 
                                    THEN hp_sum.total_abonos 
                                    ELSE COALESCE(pg.monto_total, 0) 
                                END
                            ) AS monto_pagado_real,
                            SUBSTRING_INDEX(GROUP_CONCAT(COALESCE(hp_sum.metodo_abono, pg.metodo_pago) ORDER BY pg.fecha_pago DESC, pg.id DESC SEPARATOR '||'), '||', 1) AS ultimo_metodo_pago
                        FROM pagos pg
                        LEFT JOIN (
                            SELECT 
                                pago_id, 
                                SUM(monto_pagado) AS total_abonos,
                                SUBSTRING_INDEX(GROUP_CONCAT(metodo_pago ORDER BY fecha_pago DESC, id DESC SEPARATOR '||'), '||', 1) AS metodo_abono
                            FROM historial_pagos
                            GROUP BY pago_id
                        ) hp_sum ON hp_sum.pago_id = pg.id
                        GROUP BY pg.membresia_id
                    ) pg_agg ON pg_agg.membresia_id = m.id
                ) m_det
                GROUP BY m_det.cliente_id
            ) rm ON rm.cliente_id = c.id
            LEFT JOIN (
                SELECT 
                    v.cliente_id,
                    COUNT(v.id) AS total_compras_pos,
                    SUM(COALESCE(v.monto_total, 0)) AS monto_total_ventas_pos
                FROM ventas v
                WHERE v.cliente_id IS NOT NULL
                GROUP BY v.cliente_id
            ) rv ON rv.cliente_id = c.id
            LEFT JOIN (
                SELECT 
                    ma.cliente_id,
                    COUNT(ma.id) AS total_marcajes_tabla,
                    MAX(ma.fecha_registro) AS ultima_fecha_asistencia
                FROM marcajes_asistencia ma
                GROUP BY ma.cliente_id
            ) ra ON ra.cliente_id = c.id
            WHERE {where_clause}
            ORDER BY c.id ASC;
        """)

        result = db.session.execute(sql_query, params)
        keys = list(result.keys())
        rows = [dict(zip(keys, row)) for row in result.fetchall()]

        # Convertir tipos Decimal a float para serialización limpia en JSON/Excel
        for r in rows:
            for k, v in r.items():
                if hasattr(v, "__float__") and not isinstance(v, (int, float, bool)):
                    r[k] = float(v)

        timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")

        # 5. Salida en Excel (.xlsx) si se solicita ?formato=xlsx o ?formato=excel
        if formato in ("xlsx", "excel"):
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                from openpyxl.utils import get_column_letter

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Consolidado_Clientes_BI"

                header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                thin_border = Border(
                    left=Side(style="thin", color="E2E8F0"),
                    right=Side(style="thin", color="E2E8F0"),
                    top=Side(style="thin", color="E2E8F0"),
                    bottom=Side(style="thin", color="E2E8F0"),
                )

                if keys:
                    ws.append(keys)
                    ws.row_dimensions[1].height = 28
                    for col_num in range(1, len(keys) + 1):
                        cell = ws.cell(row=1, column=col_num)
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = header_alignment

                    for row_idx, r in enumerate(rows, start=2):
                        row_vals = [r.get(k) for k in keys]
                        ws.append(row_vals)
                        for col_idx in range(1, len(keys) + 1):
                            cell = ws.cell(row=row_idx, column=col_idx)
                            cell.border = thin_border

                    # Inmovilizar panel de títulos
                    ws.freeze_panes = "A2"

                    # Auto-ajuste de ancho de columnas
                    for col in ws.columns:
                        max_len = 0
                        col_letter = get_column_letter(col[0].column)
                        for cell in col:
                            val_str = str(cell.value or "")
                            if len(val_str) > max_len:
                                max_len = len(val_str)
                        ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)

                bio = BytesIO()
                wb.save(bio)
                bio.seek(0)
                filename = f"reporte_consolidado_clientes_{timestamp_str}.xlsx"
                return Response(
                    bio.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={
                        "Content-Disposition": f'attachment; filename="{filename}"'
                    },
                )
            except ImportError:
                # Fallback automático a CSV para Excel si openpyxl falla
                formato = "csv"

        # 6. Salida en CSV compatible con Excel (UTF-8 con BOM) si se solicita ?formato=csv
        if formato == "csv":
            si = StringIO()
            si.write("\ufeff")  # BOM para que Excel abra acentos y columnas automáticamente
            cw = csv.writer(si)
            if keys:
                cw.writerow(keys)
                for r in rows:
                    cw.writerow([r.get(k, "") if r.get(k) is not None else "" for k in keys])
            filename = f"reporte_consolidado_clientes_{timestamp_str}.csv"
            return Response(
                si.getvalue(),
                mimetype="text/csv; charset=utf-8",
                headers={
                    "Content-Disposition": f'attachment; filename="{filename}"'
                },
            )

        # 7. Salida por defecto: JSON plano (óptimo para Power BI Web Connector y Google Sheets)
        return jsonify(rows), 200

    except Exception as e:
        current_app.logger.exception("Error generando el reporte consolidado BI: %s", str(e))
        return jsonify({"success": False, "error": "Error interno generando reporte", "detail": str(e)}), 500

